"""Convert the hand-written authoring workbook into the evaluation files.

Usage:  python -m src.eval_import            -> eval/questions_dev.jsonl + eval/questions_test.jsonl, then validates
        python -m src.eval_import --split dev  -> only rewrite the dev file (use while test is still being written)

Refuses to touch questions_test.jsonl once eval/LOCK.json exists.
"""
import argparse, csv, datetime, json, re, sys
from openpyxl import load_workbook
from .common import ROOT
from .eval_validate import validate

COLS = ["item_id", "split", "type", "expected_behavior", "passage_1", "passage_2", "known_duplicates",
        "p1_text", "p2_text", "question", "gold_answer", "extra_ids", "written_on", "notes"]


def ids(cell):
    return [x.strip() for x in str(cell or "").replace(";", ",").split(",") if x.strip()]


def spares():
    p = ROOT / "eval/authoring_spares.csv"
    return {r["spare_id"]: r for r in csv.DictReader(open(p, encoding="utf-8"))} if p.exists() else {}


def to_item(r, sp):
    t, beh = r["type"], (r["expected_behavior"] or "").strip().lower()
    c1 = (r["passage_1"] or "").split("|")[0].strip()
    c2 = (r["passage_2"] or "").split("|")[0].strip()
    m = re.search(r"\bS(\d)\b", str(r["notes"] or ""))
    redraw = None
    if t == "multihop" and m and f"S{m.group(1)}" in sp:
        redraw = f"S{m.group(1)}"; c1, c2 = sp[redraw]["source_chunk_1"], sp[redraw]["source_chunk_2"]
        r = dict(r, known_duplicates="")
    dup = ids(r["known_duplicates"]); extra = ids(r["extra_ids"])
    groups, distractors = [], []
    if t == "single":
        groups = [[c1] + dup + extra]
    elif t == "multihop":
        groups = [[c1] + [d for d in dup if d != c2], [c2]] + ([extra] if extra else [])
    elif t == "nearmiss":
        distractors = [c1]
        if beh == "refer":
            groups = [extra or [c1]]
    w = r["written_on"]
    w = w.date().isoformat() if isinstance(w, datetime.datetime) else (str(w).strip() if w else "")
    return {"id": r["item_id"], "split": r["split"], "type": t, "question": str(r["question"]).strip(),
            "gold_answer": str(r["gold_answer"] or "").strip(), "gold_chunk_ids": groups,
            "distractor_chunk_ids": distractors, "answerable": t in ("single", "multihop"),
            "expected_behavior": beh, "written_on": w or datetime.date.today().isoformat(),
            "source_assignment": [c for c in (c1, c2) if c], "redraw_spare": redraw}


def main():
    ap = argparse.ArgumentParser(); ap.add_argument("--split", choices=["dev", "test"]); a = ap.parse_args()
    wb = load_workbook(ROOT / "eval/authoring_workbook.xlsx", data_only=True)
    ws = wb["Items"]
    rows = [dict(zip(COLS, [c.value for c in row])) for row in ws.iter_rows(min_row=2)]
    out = {"dev": [], "test": []}; skipped = []; sp = spares()
    for r in rows:
        if not r["item_id"]: continue
        if not (r["question"] or "").strip():
            skipped.append(r["item_id"]); continue
        out[r["split"]].append(to_item(r, sp))
    for split in ([a.split] if a.split else ["dev", "test"]):
        if split == "test" and (ROOT / "eval/LOCK.json").exists():
            print("test set is LOCKED - not rewriting questions_test.jsonl"); continue
        with open(ROOT / f"eval/questions_{split}.jsonl", "w", encoding="utf-8") as f:
            for it in out[split]:
                f.write(json.dumps(it, ensure_ascii=False) + "\n")
        print(f"{split}: wrote {len(out[split])} items")
    if skipped: print("not yet written / REDRAW:", ", ".join(skipped))
    sys.exit(0 if validate(a.split) else 1)


if __name__ == "__main__":
    main()
