"""Build eval/authoring_workbook.xlsx: one row per planned evaluation item, with the full text of the
randomly assigned passage(s), so questions and gold answers can be written by hand in Excel.

Usage:  python scripts/make_authoring_workbook.py      (never overwrites an existing workbook)
Then:   python -m src.eval_import                       (workbook -> eval/questions_{dev,test}.jsonl)
"""
import csv, json, sys
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "eval/authoring_workbook.xlsx"
if OUT.exists() and "--force" not in sys.argv:
    sys.exit(f"{OUT} exists - refusing to overwrite your written questions (use --force only if it is blank)")

chunks = {json.loads(l)["chunk_id"]: json.loads(l) for l in open(ROOT / "data/processed/chunks.jsonl", encoding="utf-8")}
plan = list(csv.DictReader(open(ROOT / "eval/authoring_plan.csv", encoding="utf-8")))
dq = json.load(open(ROOT / "reports/data_quality.json", encoding="utf-8"))
# near-identical passages (the catalog repeats itself): either copy counts as correct retrieval
ALT = {}
for a, b, j in dq["near_duplicate_examples"]:
    if j >= 0.8:
        ALT.setdefault(a, []).append(b); ALT.setdefault(b, []).append(a)

def ref(cid):
    c = chunks[cid]
    return f'{cid} | {c["doc_title"]} | {" > ".join(c["heading_path"])} | p.{c["page_start"]}'

wb = Workbook()
ins = wb.active; ins.title = "READ ME"
lines = [
    "How to write the evaluation set (you write every question and gold answer yourself - no AI)",
    "",
    "1. Work in the 'Items' sheet. Write DEV rows (D01-D10) first, then TEST rows (T11-T50).",
    "2. Read the assigned passage in column H (and I for multi-hop). Write a question a graduate student would really ask, in your own words (column J).",
    "3. Gold answer (column K): the answer, taken from the passage text. Include every condition/exception the passage states.",
    "4. single: answered by passage 1 alone.  multihop: needs BOTH passages. If no natural question links them, type the next unused spare id (S1, then S2 ...) from eval/authoring_spares.csv in column N, read those two passages in the 'Chunks' sheet, and write the question for them. Never hand-pick a replacement.",
    "5. unanswerable: no passage. Write a realistic question whose answer is NOT anywhere in the corpus. Gold answer = what the system should say (e.g. 'Not in the corpus; contact the Registrar').",
    "   Check it: python -m src.eval_validate --check-absent \"your keywords\"  -> read the listed chunks; if one answers it, change the question.",
    "6. nearmiss: write a question on the passage's topic that asks for a detail the passage (and the corpus) does NOT contain.",
    "   Choose expected_behavior: 'refer' if the corpus says where the detail is published (then put that chunk id in column L), else 'abstain'.",
    "7. Column L (optional): extra correct chunk ids, comma-separated, if another passage also fully answers the question. Use the 'Chunks' sheet to search.",
    "8. Do not open or run the TEST questions through the system until they are locked. Never edit a test row after locking.",
    "9. Save, then run:  python -m src.eval_import   and   python -m src.eval_validate   (and --lock when all 50 pass).",
    "",
    "Why passages are pre-assigned: they were drawn at random (seed 13) so the test set cannot drift toward easy passages.",
]
for i, l in enumerate(lines, 1):
    ins.cell(row=i, column=1, value=l).font = Font(bold=(i == 1), size=12 if i == 1 else 11)
ins.column_dimensions["A"].width = 150

ws = wb.create_sheet("Items")
hdr = ["item_id", "split", "type", "expected_behavior", "passage_1", "passage_2", "known_duplicates",
       "PASSAGE 1 TEXT", "PASSAGE 2 TEXT", "QUESTION (you write)", "GOLD ANSWER (you write)",
       "extra / refer chunk ids", "written_on (YYYY-MM-DD)", "notes / REDRAW"]
ws.append(hdr)
for r in plan:
    t = r["type"]; c1, c2 = r["source_chunk_1"], r["source_chunk_2"]
    beh = {"single": "answer", "multihop": "answer", "unanswerable": "abstain"}.get(t, "")
    dup = ", ".join(sorted(set(ALT.get(c1, []) + ALT.get(c2, []))))
    ws.append([r["item_id"], r["split"], t, beh, ref(c1) if c1 else "", ref(c2) if c2 else "", dup,
               chunks[c1]["body"] if c1 else "(no passage: write a question the corpus cannot answer)",
               chunks[c2]["body"] if c2 else "", "", "", "", "", ""])
widths = [8, 6, 12, 12, 34, 34, 14, 70, 50, 45, 45, 22, 14, 16]
for i, w in enumerate(widths):
    ws.column_dimensions[chr(65 + i)].width = w
yellow = PatternFill("solid", fgColor="FFF2CC"); grey = PatternFill("solid", fgColor="EDEDED")
for row in ws.iter_rows(min_row=1):
    for cell in row:
        cell.alignment = Alignment(wrap_text=True, vertical="top")
for c in ws[1]: c.font = Font(bold=True); c.fill = PatternFill("solid", fgColor="D9E2F3")
for r in range(2, ws.max_row + 1):
    for col in "JKLMN": ws[f"{col}{r}"].fill = yellow
    for col in "HI": ws[f"{col}{r}"].fill = grey
    ws.row_dimensions[r].height = 150
dv = DataValidation(type="list", formula1='"answer,abstain,refer"', allow_blank=True); ws.add_data_validation(dv); dv.add(f"D2:D{ws.max_row}")
ws.freeze_panes = "E2"

cs = wb.create_sheet("Chunks")
cs.append(["chunk_id", "document", "section", "pages", "text"])
for c in chunks.values():
    cs.append([c["chunk_id"], c["doc_id"], " > ".join(c["heading_path"]), f'{c["page_start"]}-{c["page_end"]}', c["body"]])
for col, w in zip("ABCDE", [14, 9, 60, 8, 120]): cs.column_dimensions[col].width = w
for c in cs[1]: c.font = Font(bold=True)
cs.auto_filter.ref = f"A1:E{cs.max_row}"; cs.freeze_panes = "A2"
wb.save(OUT)
print("wrote", OUT, "rows:", len(plan))
